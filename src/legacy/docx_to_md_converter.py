#---------------------------------------------------------------------------------------------
#-- Скрипт для конвертації DOCX-файлів у формат Markdown (.md) за допомогою python-docx.
#-- Встановлення залежностей: pip install python-docx
#---------------------------------------------------------------------------------------------
import os
import argparse
import glob
import re # ВИПРАВЛЕНО: Додано імпорт модуля для регулярних виразів
from docx import Document # Бібліотека для роботи з DOCX
from io import StringIO

# --- КОНСТАНТИ ТА УТИЛІТИ ---

# Словник для зіставлення стилів DOCX із заголовками Markdown
HEADER_STYLES = {
    'Heading 1': '#',
    'Heading 2': '##',
    'Heading 3': '###',
    'Heading 4': '####',
    'Heading 5': '#####',
    'Heading 6': '######',
}

def get_list_prefix(paragraph):
    """Визначає, чи є параграф частиною маркованого чи нумерованого списку."""
    # У docx немає прямого легкого доступу до типу списку через API, 
    # тому доводиться покладатися на стилі. Це спрощена логіка.
    style_name = paragraph.style.name
    
    # Спрощена перевірка на стилі списків
    if 'List' in style_name or 'list' in style_name:
        # На жаль, python-docx не дає простого способу відрізнити маркований від нумерованого.
        # За замовчуванням припустимо, що це маркований список.
        return '* ' 
    return None

def convert_table_to_md(table):
    """Конвертує об'єкт таблиці python-docx у формат Markdown."""
    
    md_output = []
    
    # Витягуємо дані рядків
    data = []
    for row in table.rows:
        data.append([cell.text.strip() for cell in row.cells])
        
    if not data:
        return ""

    # Обчислюємо максимальну ширину кожного стовпця
    col_widths = [max(len(cell[i]) for cell in data) for i in range(len(data[0]))]
    
    # Додаємо верхній рядок (заголовки)
    md_output.append("| " + " | ".join(f"{data[0][i]:<{col_widths[i]}}" for i in range(len(data[0]))) + " |")
    
    # Додаємо роздільник заголовка
    separator = []
    for width in col_widths:
        separator.append("-" * width)
    md_output.append("| " + " | ".join(separator) + " |")

    # Додаємо решту рядків (тіло таблиці)
    for row_data in data[1:]:
        md_output.append("| " + " | ".join(f"{row_data[i]:<{col_widths[i]}}" for i in range(len(row_data))) + " |")
        
    return "\n".join(md_output) + "\n\n"


def extract_and_format_docx(docx_path):
    """
    Витягує текст та структуру з DOCX-файлу.
    """
    md_content_stream = StringIO()
    
    try:
        document = Document(docx_path)
    except FileNotFoundError:
        print(f"❌ Помилка: Файл не знайдено за шляхом: {docx_path}")
        return None
    except Exception as e:
        print(f"❌ Непередбачувана помилка при читанні DOCX: {e}")
        return None

    md_content_stream.write(f"\n# Документ: {os.path.basename(docx_path)}\n\n")

    # Обробка параграфів та таблиць
    for element in document.element.body:
        tag = element.tag.split('}')[-1]
        
        if tag == 'p':
            paragraph = document.paragraphs[document.element.body.index(element)]
            text = paragraph.text.strip()

            if not text:
                continue

            style_name = paragraph.style.name

            # 1. Заголовки
            if style_name in HEADER_STYLES:
                prefix = HEADER_STYLES[style_name]
                md_content_stream.write(f"{prefix} {text}\n\n")
            
            # 2. Списки (спрощена логіка)
            elif (list_prefix := get_list_prefix(paragraph)):
                 # Видаляємо маркер списку, якщо він присутній у тексті (зазвичай його там немає, 
                 # але це допомагає уникнути подвійних маркерів)
                clean_text = re.sub(r'^[\*\-\s]+', '', text).strip()
                md_content_stream.write(f"{list_prefix}{clean_text}\n")
            
            # 3. Звичайний текст
            else:
                # Вставляємо подвійний перенос рядка для нового абзацу
                md_content_stream.write(f"{text}\n\n")

        elif tag == 'tbl':
            # 4. Таблиці
            table = document.tables[document.element.body.index(element)]
            md_content_stream.write(convert_table_to_md(table))
            
    # Додаткова очистка: видаляємо зайві порожні рядки
    final_content = re.sub(r'\n\s*\n\s*\n', '\n\n', md_content_stream.getvalue())

    return final_content.strip()


def convert_docx_to_markdown(docx_path, output_dir=None):
    """
    Основна функція для конвертації DOCX у файл Markdown.
    """
    
    # 1. Визначення шляхів
    input_filename = os.path.basename(docx_path)
    base_name = os.path.splitext(input_filename)[0]
    
    if output_dir is None:
        # Збереження у тій самій директорії, що й DOCX (режим за замовчуванням)
        output_dir = os.path.dirname(docx_path) or '.' 
        
    os.makedirs(output_dir, exist_ok=True)
    
    # Використовуємо оригінальну назву, але з розширенням .md
    output_filename = f"{base_name}.md" 
    output_path = os.path.join(output_dir, output_filename)
    
    print(f"🚀 Початок конвертації: {input_filename}")
    
    # 2. Витягнення та форматування вмісту
    markdown_content = extract_and_format_docx(docx_path)
    
    if markdown_content is None:
        return
    
    # 3. Збереження результату
    try:
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(markdown_content)
        
        print(f"✅ Успішно збережено: {output_path}")
        
    except Exception as e:
        print(f"❌ Помилка збереження файлу: {e}")


def parse_arguments():
    """Обробка аргументів командного рядка."""
    parser = argparse.ArgumentParser(
        description="Конвертує один або кілька DOCX-файлів у формат Markdown (.md).",
        formatter_class=argparse.RawTextHelpFormatter
    )

    parser.add_argument(
        'input_paths',
        nargs='*',  # Нуль або більше аргументів
        help="Шлях(и) до DOCX-файлу або папки для сканування. Якщо не вказано, сканує 'data/' рекурсивно."
    )
    
    parser.add_argument(
        '-o', '--output-dir',
        type=str,
        default=None,
        help="Опціональна папка для збереження результатів. За замовчуванням: та ж папка, що й DOCX."
    )

    args = parser.parse_args()
    return args

if __name__ == "__main__":
    
    args = parse_arguments()
    files_to_process = []
    
    # Обробка вхідних шляхів (файли або папки)
    if not args.input_paths:
        print("🔍 Режим: Аргументи не вказано. Рекурсивний пошук DOCX у папці 'data/'...")
        # Використовуємо glob для рекурсивного пошуку всіх DOCX у data/
        search_pattern = 'data/**/*.docx' 
        files_to_process = glob.glob(search_pattern, recursive=True)
        
    else:
        # Логіка обробки наданих аргументів (якщо вони є)
        for path in args.input_paths:
            if os.path.isdir(path):
                # Якщо це папка, шукаємо всі DOCX-файли в ній
                for file in os.listdir(path):
                    if file.lower().endswith('.docx'):
                        files_to_process.append(os.path.join(path, file))
            elif os.path.isfile(path) and path.lower().endswith('.docx'):
                # Якщо це єдиний DOCX-файл
                files_to_process.append(path)
            else:
                print(f"⚠️ Пропущено: Шлях '{path}' не є DOCX-файлом або папкою з DOCX.")


    if not files_to_process:
        print("🛑 Не знайдено DOCX-файлів для обробки. Переконайтеся, що вони існують у 'data/'")
        exit(0)
        
    print(f"\n✅ Знайдено {len(files_to_process)} DOCX-файлів для конвертації.")
    
    for docx_file in files_to_process:
        # Зберігаємо файл у тій самій директорії, якщо -o не вказано
        convert_docx_to_markdown(docx_file, args.output_dir)

    print("\n" + "=" * 50)
    print("🎉 Конвертація завершена.")
    print("========================================")