"""
DOCX to Markdown converter service
Migrated from: src/legacy/docx_to_md_converter.py
"""
import re
import structlog
from typing import Optional
from io import BytesIO
from docx import Document


logger = structlog.get_logger()


# Mapping of DOCX styles to Markdown headers
HEADER_STYLES = {
    'Heading 1': '#',
    'Heading 2': '##',
    'Heading 3': '###',
    'Heading 4': '####',
    'Heading 5': '#####',
    'Heading 6': '######',
}


class DOCXConverter:
    """
    Service for converting DOCX documents to Markdown format
    """
    
    def __init__(self):
        pass
    
    def _get_list_prefix(self, paragraph) -> Optional[str]:
        """Determine if paragraph is part of a list"""
        style_name = paragraph.style.name
        
        # Simplified check for list styles
        if 'List' in style_name or 'list' in style_name:
            return '* '
        return None
    
    def _convert_table_to_md(self, table) -> str:
        """Convert DOCX table to Markdown table"""
        md_output = []
        
        # Extract row data
        data = []
        for row in table.rows:
            data.append([cell.text.strip() for cell in row.cells])
        
        if not data:
            return ""
        
        # Calculate maximum width for each column
        col_widths = [max(len(str(cell[i])) for cell in data) for i in range(len(data[0]))]
        
        # Add header row
        md_output.append("| " + " | ".join(f"{data[0][i]:<{col_widths[i]}}" for i in range(len(data[0]))) + " |")
        
        # Add separator
        separator = ["-" * width for width in col_widths]
        md_output.append("| " + " | ".join(separator) + " |")
        
        # Add body rows
        for row_data in data[1:]:
            md_output.append("| " + " | ".join(f"{str(row_data[i]):<{col_widths[i]}}" for i in range(len(row_data))) + " |")
        
        return "\n".join(md_output) + "\n\n"
    
    async def convert_to_markdown(self, docx_content: bytes, filename: str) -> Optional[str]:
        """
        Convert DOCX bytes to Markdown text
        
        Args:
            docx_content: DOCX file content as bytes
            filename: Original filename for reference
            
        Returns:
            Markdown text or None if conversion fails
        """
        logger.info("convert_docx_to_markdown", filename=filename, size=len(docx_content))
        
        try:
            # Use BytesIO to read DOCX from memory
            docx_bytes = BytesIO(docx_content)
            document = Document(docx_bytes)
            
            md_lines = []
            md_lines.append(f"\n# Документ: {filename}\n\n")
            
            # Process paragraphs and tables
            for element in document.element.body:
                tag = element.tag.split('}')[-1]
                
                if tag == 'p':
                    # Find paragraph by index
                    para_index = list(document.element.body).index(element)
                    paragraph = document.paragraphs[para_index]
                    text = paragraph.text.strip()
                    
                    if not text:
                        continue
                    
                    style_name = paragraph.style.name
                    
                    # 1. Headers
                    if style_name in HEADER_STYLES:
                        prefix = HEADER_STYLES[style_name]
                        md_lines.append(f"{prefix} {text}\n\n")
                    
                    # 2. Lists
                    elif (list_prefix := self._get_list_prefix(paragraph)):
                        clean_text = re.sub(r'^[\*\-\s]+', '', text).strip()
                        md_lines.append(f"{list_prefix}{clean_text}\n")
                    
                    # 3. Regular text
                    else:
                        md_lines.append(f"{text}\n\n")
                
                elif tag == 'tbl':
                    # 4. Tables
                    tbl_index = list(document.element.body).index(element)
                    # Find corresponding table
                    table_counter = 0
                    for el in list(document.element.body)[:tbl_index + 1]:
                        if el.tag.split('}')[-1] == 'tbl':
                            table_counter += 1
                    
                    if table_counter <= len(document.tables):
                        table = document.tables[table_counter - 1]
                        md_lines.append(self._convert_table_to_md(table))
            
            # Join and clean up
            markdown_content = "".join(md_lines)
            markdown_content = re.sub(r'\n\s*\n\s*\n', '\n\n', markdown_content)
            
            logger.info(
                "docx_converted",
                filename=filename,
                output_length=len(markdown_content)
            )
            
            return markdown_content.strip()
            
        except Exception as e:
            logger.error(
                "docx_conversion_failed",
                filename=filename,
                error=str(e)
            )
            return None
